"""F9 step 2 and doc 09 «Legales» (decision 095): legal documents with the
real data of the file — never fixed fake text (reference defects #3, #16,
#33). Each one is built once as plain paragraphs and rendered as the
issued PDF (kept with its sha256) and as an editable DOCX for the lawyer
or the notary.

The texts are models reviewed against Dominican law (decision 127,
25/10/2026): Ley 45-20 on garantías mobiliarias (the pledge is a
«garantía mobiliaria sin desposesión» registered in the SEGM, with the
extrajudicial execution of its arts. 87 and 94-104), Ley 108-05 and the
SCJ criterion for the mortgage (acto bajo firma privada with legalized
signatures, registered at the Registro de Títulos), Ley 140-15 for the
notarial note (amounts and dates in words and figures, identification of
the parties, reading before signing; art. 545 CPC for its executory
force), art. 583 CPC for the payment order, Ley 358-05 art. 53 for the
credit disclosure, and Ley 172-13 art. 5 for the written consent to the
credit bureaus. Each company may still adjust them with its lawyer (the
DOCX is editable).
"""

import io
from dataclasses import dataclass, field
from decimal import Decimal

from django.utils import timezone

from apps.collections.models import LegalCase, LegalProfessional
from apps.core.text import plural
from apps.documents.models import GeneratedDocument
from apps.documents.services import amount_in_words, date_in_words, issue
from apps.lending.models import Loan

ZERO = Decimal("0.00")


class LegalDocumentError(Exception):
    def __init__(self, message: str, code: str = "invalid"):
        super().__init__(message)
        self.message = message
        self.code = code


@dataclass
class LegalText:
    title: str
    paragraphs: list[str]
    signatures: list[tuple[str, str]] = field(default_factory=list)  # (name, role)


def _customer_line(loan: Loan) -> str:
    c = loan.customer
    parts = [f"{c.display_name}, dominicano(a), mayor de edad"]
    if c.marital_status:
        parts.append(c.marital_status.lower())
    if c.occupation:
        parts.append(c.occupation.lower())
    parts.append(
        f"portador(a) de la cédula de identidad y electoral núm. {_doc(c.document_number)}"
    )
    if c.address_line:
        parts.append(f"domiciliado(a) en {c.address_line}")
    return ", ".join(parts)


def _doc(value: str) -> str:
    digits = "".join(ch for ch in value or "" if ch.isdigit())
    if len(digits) == 11:
        return f"{digits[:3]}-{digits[3:10]}-{digits[10]}"
    return value or "—"


def _company_line(loan: Loan) -> str:
    company = loan.company
    line = company.legal_name
    if company.rnc:
        line += f", RNC núm. {company.rnc}"
    if company.address:
        line += f", con domicilio en {company.address}"
    return line


def _professional_line(person: LegalProfessional, role: str) -> str:
    line = f"{person.name}, {role}"
    if person.document:
        line += f", cédula núm. {_doc(person.document)}"
    if person.license:
        line += f", matrícula núm. {person.license}"
    if person.address:
        line += f", con estudio en {person.address}"
    return line


def exigible(loan: Loan) -> Decimal:
    """What is due today: the overdue installments with their late fees."""
    from apps.payments import portfolio

    today = timezone.localdate()
    return sum(
        (
            line.total
            for line in portfolio.position(loan, today)
            if line.installment.due_date < today
        ),
        ZERO,
    )


def _money(amount: Decimal) -> str:
    return f"{amount_in_words(amount)} (RD$ {amount:,.2f})"


def _pct(value: Decimal) -> str:
    text = f"{value:.4f}".rstrip("0").rstrip(".")
    return f"{text} %"


def _terms_line(loan: Loan) -> str:
    """The credit conditions a Dominican consumer must see (Ley 358-05,
    art. 53): rate, late fee and the costs, as the pagaré and the schedule
    state them."""
    from apps.documents.services import FREQUENCY_LABELS
    from apps.payments import portfolio
    from engine.rates import RateUnit, monthly_equivalent

    terms = portfolio.current_terms(loan)
    singular, _plural = FREQUENCY_LABELS[terms.frequency]
    monthly = monthly_equivalent(terms.rate, RateUnit(terms.rate_unit), terms.frequency) * 100
    grace = plural(terms.late_fee_grace_days, "día")
    if terms.late_fee_method == "fixed":
        fee = f"un cargo fijo de RD$ {terms.late_fee_fixed_amount or ZERO:,.2f} por cuota"
    elif terms.late_fee_method == "full_period":
        fee = (
            f"una mora del {_pct(terms.late_fee_monthly_rate * 100)} por cada período completo "
            "de atraso sobre lo vencido de la cuota"
        )
    else:
        fee = (
            f"una mora del {_pct(terms.late_fee_monthly_rate * 100)} mensual, prorrateada por "
            "día, sobre lo vencido de la cuota"
        )
    return (
        f"con un interés del {_pct(terms.periodic_rate * 100)} por período {singular} "
        f"(equivalente al {_pct(monthly)} mensual) y, pasados {grace} de gracia desde el "
        f"vencimiento de una cuota, {fee}, sin que la mora de una cuota supere el "
        f"{_pct(terms.late_fee_cap_pct * 100)} de su saldo vencido. El costo total del crédito, "
        "los cargos y el costo anual efectivo constan en el pagaré y en la tabla de pagos, que "
        "el deudor declara haber recibido"
    )


CONSENT = (
    "EL DEUDOR autoriza a EL ACREEDOR a tratar sus datos personales para la gestión de este "
    "crédito y, de manera libre, expresa y por escrito, a consultar y reportar su historial a "
    "las sociedades de información crediticia, conforme a la Ley 172-13 sobre protección de "
    "datos de carácter personal."
)


# --- The texts -------------------------------------------------------------------------------


def demand_notice(case: LegalCase) -> LegalText:
    """Acto de intimación de pago: only with real arrears, for the exact
    exigible amount (defect #33)."""
    loan = case.loan
    owed = exigible(loan)
    if owed <= 0 or loan.days_past_due <= 0:
        raise LegalDocumentError("Sin atraso real no se emite la intimación.", "not_late")
    today = timezone.localdate()
    return LegalText(
        title="ACTO DE INTIMACIÓN DE PAGO",
        paragraphs=[
            f"En la ciudad de su emisión, a los {date_in_words(today)}.",
            f"A requerimiento de {_company_line(loan)}, quien tiene como abogado apoderado al "
            f"{_professional_line(case.lawyer, 'abogado de los tribunales de la República')}.",
            f"Se INTIMA a {_customer_line(loan)}, deudor(a) del préstamo núm. {loan.number} "
            f"desembolsado el {date_in_words(loan.disbursed_on)} por "
            f"{_money(loan.principal)}.",
            f"Para que en el plazo de cinco (5) días francos pague la suma de {_money(owed)}, "
            f"correspondiente a las cuotas vencidas y su mora al {date_in_words(today)}, con "
            f"{plural(loan.days_past_due, 'día')} de atraso, en las oficinas del requeriente o "
            "por los medios de pago que este le indique, más lo que venza hasta el pago.",
            "Advertido de que, a falta de pago en el plazo indicado, el requeriente procederá a "
            "ejercer las vías de derecho que le acuerda la ley para el cobro de su crédito, "
            "incluida la ejecución de las garantías, con los gastos a cargo del deudor.",
        ],
        signatures=[(case.lawyer.name, "Abogado apoderado")],
    )


def payment_order(case: LegalCase) -> LegalText:
    """Mandamiento de pago, notified by the bailiff of the case."""
    loan = case.loan
    if case.bailiff is None:
        raise LegalDocumentError("Asigna el alguacil del caso.", "bailiff")
    owed = exigible(loan)
    if owed <= 0:
        raise LegalDocumentError("Sin atraso real no hay mandamiento de pago.", "not_late")
    today = timezone.localdate()
    return LegalText(
        title="MANDAMIENTO DE PAGO",
        paragraphs=[
            f"En la ciudad de su emisión, a los {date_in_words(today)}, actuando a "
            f"requerimiento de {_company_line(loan)}, representada por su abogado apoderado "
            f"{_professional_line(case.lawyer, 'abogado')}.",
            f"Yo, {_professional_line(case.bailiff, 'alguacil')}, notifico a "
            f"{_customer_line(loan)} el presente MANDAMIENTO DE PAGO.",
            f"Por el cual se le requiere pagar dentro de un (1) día franco la suma de "
            f"{_money(owed)}, adeudada en virtud del préstamo núm. {loan.number}, más los "
            "intereses y gastos que se generen hasta el pago total, en virtud del título "
            "ejecutorio del que se le notifica copia en cabeza del presente acto (artículo 583 "
            "del Código de Procedimiento Civil).",
            "Advertido de que, vencido ese plazo sin haber pagado, se procederá al embargo "
            "ejecutivo de sus bienes muebles conforme a los artículos 583 y siguientes del Código "
            "de Procedimiento Civil, con los gastos a su cargo.",
        ],
        signatures=[(case.bailiff.name, "Alguacil")],
    )


def voluntary_surrender(case: LegalCase) -> LegalText:
    """Acto de entrega voluntaria: the pledged asset, described from
    Garantías (defect #16), before two witnesses of the lawyer."""
    loan = case.loan
    pledged = [
        item.collateral
        for item in loan.collaterals.select_related("collateral__vehicle", "collateral__property")
    ]
    if not pledged:
        raise LegalDocumentError("El préstamo no tiene bienes en garantía.", "collateral")
    descriptions = [_asset(asset) for asset in pledged]
    witnesses = list(case.lawyer.witnesses.filter(is_active=True)[:2])
    today = timezone.localdate()
    return LegalText(
        title="ACTO DE ENTREGA VOLUNTARIA",
        paragraphs=[
            f"En la ciudad de su emisión, a los {date_in_words(today)}.",
            f"{_customer_line(loan)}, en su calidad de deudor(a) del préstamo núm. {loan.number} "
            f"otorgado por {_company_line(loan)}.",
            "DECLARA que entrega de manera libre y voluntaria, en garantía del pago de lo "
            "adeudado, el (los) bien(es) siguiente(s):",
            *[f"— {text}" for text in descriptions],
            "El bien se recibe en el estado en que se encuentra, conforme al procedimiento de "
            "ejecución convenido en el contrato de garantía (Ley 45-20 sobre garantías "
            "mobiliarias, artículos 87 y 88). Su venta se hará según los artículos 94 al 104 de "
            "esa ley: el producto se aplicará a los gastos, los intereses y el saldo del "
            "préstamo, y el excedente, si lo hubiere, se devolverá al deudor en el plazo de ley; "
            "si el producto no alcanza, el deudor sigue obligado por el saldo insoluto.",
            "El deudor conserva el derecho de pagar la totalidad de lo adeudado y los gastos "
            "antes de la venta, con lo cual recuperará el bien (artículos 97 y 121).",
            *(
                [
                    "En presencia de los testigos "
                    + " y ".join(_professional_line(w, "testigo") for w in witnesses)
                    + "."
                ]
                if witnesses
                else []
            ),
        ],
        signatures=[
            (loan.customer.display_name, "Deudor(a)"),
            (case.lawyer.name, "Por el acreedor"),
            *[(w.name, "Testigo") for w in witnesses],
        ],
    )


def _asset(asset) -> str:
    text = f"{asset.get_type_display()}: {asset.description}"
    vehicle = getattr(asset, "vehicle", None)
    if vehicle is not None:
        bits = [
            f"marca {vehicle.brand}",
            f"modelo {vehicle.model}",
            f"año {vehicle.year}" if vehicle.year else "",
            f"color {vehicle.color}" if vehicle.color else "",
            f"placa {vehicle.plate}" if vehicle.plate else "",
            f"chasis {vehicle.chassis}" if vehicle.chassis else "",
            f"matrícula {vehicle.registration}" if vehicle.registration else "",
        ]
        text += " (" + ", ".join(bit for bit in bits if bit) + ")"
    prop = getattr(asset, "property", None)
    if prop is not None:
        bits = [
            f"certificado de título {prop.title_number}" if prop.title_number else "",
            f"designación catastral {prop.cadastral_designation}"
            if prop.cadastral_designation
            else "",
            f"ubicado en {prop.location}" if prop.location else "",
        ]
        text += " (" + ", ".join(bit for bit in bits if bit) + ")"
    return text


def notarial_note(case: LegalCase) -> LegalText:
    """Pagaré notarial with the real data of the file (defect #3)."""
    loan = case.loan
    if case.notary is None:
        raise LegalDocumentError("Asigna el notario del caso.", "notary")
    from apps.payments import portfolio

    terms = portfolio.current_terms(loan)
    today = timezone.localdate()
    return LegalText(
        title="PAGARÉ NOTARIAL",
        paragraphs=[
            f"En la ciudad de su emisión, a los {date_in_words(today)}, por ante mí, "
            f"{_professional_line(case.notary, 'notario público')}, comparecen:",
            f"De una parte, {_company_line(loan)}, en lo adelante EL ACREEDOR.",
            f"De la otra parte, {_customer_line(loan)}, en lo adelante EL DEUDOR.",
            f"EL DEUDOR reconoce adeudar a EL ACREEDOR la suma de {_money(loan.principal)}, "
            f"recibida en virtud del préstamo núm. {loan.number} del "
            f"{date_in_words(loan.disbursed_on)}, pagadera en "
            f"{plural(terms.installments, 'cuota')} según la tabla de pagos que forma parte "
            f"de este acto, {_terms_line(loan)}.",
            "En caso de falta de pago de una o más cuotas a su vencimiento, EL ACREEDOR podrá "
            "exigir el pago total de lo adeudado, con su mora. EL DEUDOR puede pagar por "
            "anticipado, en todo o en parte, con la reducción proporcional de los intereses no "
            "devengados (Ley 358-05, artículo 55).",
            "La primera copia de este acto vale título ejecutorio conforme al artículo 545 del "
            "Código de Procedimiento Civil, y EL DEUDOR elige domicilio en la dirección indicada "
            "para toda notificación.",
            CONSENT,
            "Leído el presente acto en alta voz a los comparecientes, quienes declaran haberlo "
            "entendido, lo firman junto a mí, notario, que doy fe.",
        ],
        signatures=[
            (loan.customer.display_name, "El deudor"),
            ("Por el acreedor", "El acreedor"),
            (case.notary.name, "Notario público"),
        ],
    )


# --- Contracts of a loan with collateral (doc 09 «Legales») -----------------------------------


def _guarantors_line(loan: Loan) -> str:
    names = [
        f"{g.customer.display_name} (cédula núm. {_doc(g.customer.document_number)})"
        for g in loan.guarantors.select_related("customer")
    ]
    return ", ".join(names)


def _secured_amount(loan: Loan) -> Decimal:
    """The maximum secured sum (Ley 45-20 art. 16; the Registro de Títulos asks
    for it too): the capital plus every scheduled interest, rounded up to the
    next thousand so late fees and costs fit."""
    from apps.payments import portfolio

    scheduled = sum(
        (row.principal_due + row.interest_due for row in portfolio.current_installments(loan)),
        ZERO,
    )
    base = max(scheduled, loan.principal)
    return (base / 1000).to_integral_value(rounding="ROUND_CEILING") * 1000


def _contract(
    loan: Loan, *, title: str, kind_text: str, assets: list, clauses: list[str]
) -> LegalText:
    from apps.payments import portfolio

    terms = portfolio.current_terms(loan)
    guarantors = _guarantors_line(loan)
    paragraphs = [
        f"Entre {_company_line(loan)}, en lo adelante EL ACREEDOR, y {_customer_line(loan)}, en "
        "lo adelante EL DEUDOR, se ha convenido lo siguiente:",
        f"PRIMERO: EL ACREEDOR otorga a EL DEUDOR un préstamo por {_money(loan.principal)}, "
        f"núm. {loan.number}, desembolsado el {date_in_words(loan.disbursed_on)}, pagadero en "
        f"{plural(terms.installments, 'cuota')} según la tabla de pagos anexa, que forma parte "
        f"de este contrato, {_terms_line(loan)}.",
        f"SEGUNDO: En garantía del pago de ese préstamo, de sus intereses, mora y gastos de "
        f"cobro, hasta la suma máxima de {_money(_secured_amount(loan))}, EL DEUDOR otorga "
        f"{kind_text} sobre el (los) bien(es) siguiente(s), de su propiedad, libre(s) de otros "
        "gravámenes salvo los que declara aquí:",
        *[f"— {_asset(asset)}" for asset in assets],
        "TERCERO: EL DEUDOR se obliga a conservar el bien en buen estado, a no venderlo, gravarlo "
        "ni trasladarlo sin la autorización escrita de EL ACREEDOR, a mantenerlo asegurado si "
        "así se le requiere y a permitir su inspección.",
        "CUARTO: La falta de pago de una o más cuotas a su vencimiento, pasados los días de "
        "gracia, hace exigible la totalidad de lo adeudado, con su mora, y faculta a EL ACREEDOR "
        "a ejecutar la garantía. EL DEUDOR puede pagar por anticipado, en todo o en parte, con "
        "la reducción proporcional de los intereses no devengados (Ley 358-05, artículo 55).",
        *clauses,
    ]
    if guarantors:
        paragraphs.append(
            f"Intervienen como fiadores solidarios {guarantors}, quienes responden con EL DEUDOR "
            "por todas las obligaciones de este contrato y renuncian a los beneficios de "
            "excusión y de división (artículos 2021 y 2026 del Código Civil)."
        )
    paragraphs.append(CONSENT)
    paragraphs.append(
        "Para todo lo no previsto, las partes se remiten al derecho común y eligen domicilio en "
        "las direcciones indicadas, donde valen todas las notificaciones. Hecho y firmado en "
        f"dos originales, uno para cada parte, a los {date_in_words(loan.disbursed_on)}."
    )
    return LegalText(
        title=title,
        paragraphs=paragraphs,
        signatures=[
            ("Por el acreedor", _company_line(loan).split(",")[0]),
            (loan.customer.display_name, "El deudor"),
            *[
                (g.customer.display_name, "Fiador solidario")
                for g in loan.guarantors.select_related("customer")
            ],
        ],
    )


def _assets(loan: Loan, kind: str) -> list:
    return [
        item.collateral
        for item in loan.collaterals.select_related("collateral__vehicle", "collateral__property")
        if item.collateral.type == kind
    ]


def pledge_contract(loan: Loan) -> LegalText:
    """Vehicle: a garantía mobiliaria sin desposesión (Ley 45-20). Its art. 16
    asks for the parties, the maximum secured sum, a domicile, the asset, the
    obligation, the express authorization to register it in the SEGM and the
    agreement on the extrajudicial execution; art. 46 publicizes it in the SEGM."""
    assets = _assets(loan, "vehicle")
    if not assets:
        raise LegalDocumentError("El préstamo no tiene un vehículo en garantía.", "collateral")
    return _contract(
        loan,
        title="CONTRATO DE PRÉSTAMO CON GARANTÍA MOBILIARIA SOBRE VEHÍCULO",
        kind_text="garantía mobiliaria sin desposesión (Ley 45-20)",
        assets=assets,
        clauses=[
            "QUINTO: EL DEUDOR autoriza expresamente a EL ACREEDOR a inscribir esta garantía en "
            "el Sistema Electrónico de Garantías Mobiliarias (SEGM), con lo que queda oponible "
            "a terceros (artículos 16 y 46 de la Ley 45-20); los gastos de inscripción, de "
            "renovación y de cancelación son por cuenta de EL DEUDOR. Pagado el préstamo, EL "
            "ACREEDOR cancelará la inscripción en el plazo de ley.",
            "SEXTO: Las partes convienen la ejecución extrajudicial de la garantía (artículos "
            "87 y 94 al 104 de la Ley 45-20): ante el incumplimiento, EL ACREEDOR inscribirá el "
            "formulario de ejecución, EL DEUDOR entregará el bien en el plazo de tres (3) días "
            "francos desde el requerimiento o lo hará el alguacil, y la venta se hará por el "
            "ejecutor que designe EL ACREEDOR, con aviso previo. EL DEUDOR puede detener la "
            "ejecución pagando la totalidad de lo adeudado y los gastos antes de la venta; el "
            "excedente de la venta se le devuelve y el saldo insoluto, si lo hay, sigue a su "
            "cargo.",
        ],
    )


def mortgage_contract(loan: Loan) -> LegalText:
    """Real estate: hipoteca convencional under Ley 108-05, by acto bajo firma
    privada with the signatures legalized by a notary (criterion of the SCJ,
    Primera Sala) and registered at the Registro de Títulos of the property."""
    assets = _assets(loan, "property")
    if not assets:
        raise LegalDocumentError("El préstamo no tiene un inmueble en garantía.", "collateral")
    return _contract(
        loan,
        title="CONTRATO DE PRÉSTAMO CON GARANTÍA HIPOTECARIA",
        kind_text="hipoteca convencional en primer rango",
        assets=assets,
        clauses=[
            "QUINTO: Esta hipoteca se otorga conforme a la Ley 108-05 de Registro Inmobiliario. "
            "Las firmas de este acto serán legalizadas por notario público y EL ACREEDOR lo "
            "inscribirá en el Registro de Títulos de la jurisdicción del inmueble, lo que EL "
            "DEUDOR autoriza expresamente; los impuestos y gastos de inscripción, de "
            "radiación y de ejecución son por cuenta de EL DEUDOR. Saldado el préstamo, EL "
            "ACREEDOR entregará el acto de radiación.",
            "SEXTO: Ante el incumplimiento, EL ACREEDOR podrá perseguir la venta del inmueble "
            "por el procedimiento de embargo inmobiliario que establece la ley, sin perjuicio de "
            "las demás vías de derecho. EL DEUDOR declara que el inmueble no está afectado por "
            "otra hipoteca, embargo ni oposición, salvo lo que declara en este acto.",
        ],
    )


CONTRACTS = {
    GeneratedDocument.Kind.PLEDGE_CONTRACT: (pledge_contract, "vehicle", "CPR"),
    GeneratedDocument.Kind.MORTGAGE_CONTRACT: (mortgage_contract, "property", "CHI"),
}


def contract_kinds(loan: Loan) -> list[str]:
    types = set(loan.collaterals.values_list("collateral__type", flat=True))
    return [kind for kind, (_b, asset_type, _s) in CONTRACTS.items() if asset_type in types]


def issue_contract(loan: Loan, kind: str, *, issued_by) -> GeneratedDocument:
    builder, _type, suffix = CONTRACTS[kind]
    return issue(
        company=loan.company,
        kind=kind,
        target=loan,
        folio=f"{loan.number}-{suffix}",
        template="documents/legal.html",
        context={"text": builder(loan), "branch": loan.branch},
        issued_by=issued_by,
    )


BUILDERS = {
    GeneratedDocument.Kind.DEMAND_NOTICE: demand_notice,
    GeneratedDocument.Kind.PAYMENT_ORDER: payment_order,
    GeneratedDocument.Kind.VOLUNTARY_SURRENDER: voluntary_surrender,
    GeneratedDocument.Kind.NOTARIAL_NOTE: notarial_note,
}


def build(case: LegalCase, kind: str) -> LegalText:
    builder = BUILDERS.get(kind)
    if builder is None:
        raise LegalDocumentError("Documento legal desconocido.", "kind")
    return builder(case)


def issue_legal(case: LegalCase, kind: str, *, issued_by) -> GeneratedDocument:
    """The PDF, kept with its sha256 (reprints are identical)."""
    from apps.core.models import Company

    text = build(case, kind)
    stamp = timezone.localtime().strftime("%Y%m%d%H%M%S")
    return issue(
        company=Company.objects.get(pk=case.company_id),
        kind=kind,
        target=case,
        folio=f"{case.number}-{kind[:3].upper()}-{stamp}",
        template="documents/legal.html",
        context={"text": text, "branch": case.loan.branch},
        issued_by=issued_by,
    )


def docx(case: LegalCase, kind: str) -> bytes:
    """The editable version for the lawyer or the notary."""
    return render_docx(build(case, kind), case.company_id)


def contract_docx(loan: Loan, kind: str) -> bytes:
    if kind not in CONTRACTS:
        raise LegalDocumentError("Contrato desconocido.", "kind")
    return render_docx(CONTRACTS[kind][0](loan), loan.company_id)


def render_docx(text: LegalText, company_id) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    from apps.core.models import Company

    company = Company.objects.get(pk=company_id)
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    header = document.add_paragraph(company.legal_name)
    header.runs[0].bold = True
    document.add_paragraph(
        "Modelo de Cobradía revisado conforme a la ley dominicana; la empresa puede ajustarlo "
        "con su abogado."
    ).runs[0].italic = True
    heading = document.add_paragraph(text.title)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.runs[0].bold = True
    for paragraph in text.paragraphs:
        body = document.add_paragraph(paragraph)
        body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    document.add_paragraph()
    for name, role in text.signatures:
        document.add_paragraph("\n\n______________________________")
        document.add_paragraph(f"{name}\n{role}")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
