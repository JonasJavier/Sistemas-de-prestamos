"""Decision 102: what a document may load, and spreadsheets without formulas."""

import io
from pathlib import Path

import pytest
from django.conf import settings
from django.core.files.base import ContentFile
from openpyxl import load_workbook
from PIL import Image

from apps.core.tenancy import company_context
from apps.documents.services import check_resource, issue_account_statement
from apps.payments.tests.test_payments import make_loan
from apps.reports.catalog import Column, Table
from apps.reports.render import to_xlsx


def test_documents_load_embedded_data_and_media_only(tmp_path) -> None:
    check_resource("data:text/plain;base64,aG9sYQ==")
    check_resource((Path(settings.MEDIA_ROOT) / "logos" / "a.png").resolve().as_uri())
    for url in (
        "http://169.254.169.254/latest/meta-data/",
        "https://example.com/logo.png",
        Path(settings.BASE_DIR / "manage.py").resolve().as_uri(),
        (tmp_path / "secret.txt").as_uri(),
    ):
        with pytest.raises(ValueError, match="no permitido"):
            check_resource(url)


def test_the_company_logo_still_reaches_its_documents(tenant) -> None:
    buffer = io.BytesIO()
    Image.new("RGB", (40, 20), "#0e6b4f").save(buffer, format="PNG")
    company = tenant.company
    company.logo.save("logo.png", ContentFile(buffer.getvalue()), save=True)
    loan = make_loan(tenant, first_due_in=0)
    with company_context(company.id):
        document = issue_account_statement(loan, issued_by=None)
        pdf = document.file.read()
    assert b"/Subtype /Image" in pdf


def test_spreadsheet_text_never_becomes_a_formula(tenant) -> None:
    table = Table(
        title="Clientes",
        filters_line="Todos",
        columns=[Column("name", "Nombre"), Column("note", "Nota")],
        rows=[{"name": '=HYPERLINK("http://x.test","clic")', "note": "=1+1"}],
    )
    company = tenant.company
    company.legal_name = "=CMD()"
    sheet = load_workbook(io.BytesIO(to_xlsx(table, company=company, user=tenant.users["gerente"])))
    cells = [cell for row in sheet.active.iter_rows() for cell in row if cell.value]
    assert all(cell.data_type != "f" for cell in cells)
    assert '=HYPERLINK("http://x.test","clic")' in [cell.value for cell in cells]
