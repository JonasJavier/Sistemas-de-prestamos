"""D8 (decision 132): the copy of the data, the cancellation and the purge."""

import datetime as dt
import io

from django.core import mail
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.conftest import open_cash
from apps.core.models import Company
from apps.core.tenancy import company_context
from apps.customers.models import Customer
from apps.payments.models import Payment
from apps.payments.tests.test_payments import make_loan, pay
from apps.platform import cancellation

PASSWORD = "una-clave-larga-123"  # the tenant fixture's, no 2FA app in tests
EXPORT = "/api/v1/mi-empresa/datos/"
CANCEL = "/api/v1/mi-empresa/cancelar/"


def test_the_manager_downloads_the_whole_company_in_excel(tenant) -> None:
    open_cash(tenant)
    loan = make_loan(tenant, first_due_in=7)
    pay(tenant, loan, "350")
    response = tenant.client("gerente").get(EXPORT)
    assert response.status_code == 200
    assert response["Content-Disposition"].startswith('attachment; filename="cobradia-')
    book = load_workbook(io.BytesIO(response.content), read_only=True)
    names = book.sheetnames
    assert names[:2] == ["Léeme", "Usuarios"]
    for sheet in ("Clientes", "Préstamos", "Cuotas", "Pagos", "Cajas", "Asientos", "Auditoría"):
        assert sheet in names
    loans = list(book["Préstamos"].iter_rows(values_only=True))
    assert "number" in loans[0] and loans[1][loans[0].index("number")] == loan.number
    payments = list(book["Pagos"].iter_rows(values_only=True))
    assert len(payments) == 2 and "350.00" in payments[1]
    users = list(book["Usuarios"].iter_rows(values_only=True))
    assert any("gerente@uno.do" in row for row in users[1:])
    # Only the management roles.
    assert tenant.client("cajero").get(EXPORT).status_code == 403


def test_cancelling_needs_the_reason_and_the_identity(tenant) -> None:
    manager = tenant.client("gerente")
    assert manager.post(CANCEL, {"reason": "", "second_factor": PASSWORD}).status_code == 400
    wrong = manager.post(CANCEL, {"reason": "Cierro el negocio", "second_factor": "nope"})
    assert wrong.status_code == 400 and "second_factor" in wrong.json()
    cashier = tenant.client("cajero")
    assert cashier.post(CANCEL, {"reason": "x", "second_factor": PASSWORD}).status_code == 403
    tenant.company.refresh_from_db()
    assert tenant.company.status == Company.Status.TRIAL


def test_cancelling_closes_the_company_read_only_and_tells_everyone(
    tenant, django_capture_on_commit_callbacks
) -> None:
    User.objects.create_superuser("jonas@cobradia.com", "una-clave-larga-123")
    manager = tenant.client("gerente")
    mail.outbox.clear()
    with django_capture_on_commit_callbacks(execute=True):
        response = manager.post(CANCEL, {"reason": "Cierro el negocio", "second_factor": PASSWORD})
    assert response.status_code == 200, response.content
    data = response.json()
    assert data["status"] == "closed" and data["purge_on"]
    expected = (timezone.localdate() + dt.timedelta(days=90)).isoformat()
    assert data["purge_on"] == expected
    assert "Cancelada por Gina" in data["status_reason"]

    subjects = sorted(m.subject for m in mail.outbox)
    assert any(s.startswith("Cancelaste tu suscripción") for s in subjects)
    assert any(s.startswith("Cancelación en Cobradía") for s in subjects)

    # Read-only from now on, but the copy still downloads.
    with company_context(tenant.company.id):
        customer = Customer.objects.first()
    blocked = manager.patch(f"/api/v1/clientes/{customer.pk}/", {"phone": "8095550199"})
    assert blocked.status_code == 403
    assert manager.get(EXPORT).status_code == 200
    # Twice is an error, not a second closing.
    again = manager.post(CANCEL, {"reason": "otra vez", "second_factor": PASSWORD})
    assert again.status_code == 400


def test_the_purge_deletes_the_company_after_90_days(tenant) -> None:
    open_cash(tenant)
    loan = make_loan(tenant, first_due_in=7)
    pay(tenant, loan, "350")
    company = tenant.company
    cancellation.cancel_company(company=company, by=tenant.users["gerente"], reason="Me voy")
    company_id = company.id

    # Not yet: 89 days.
    Company.objects.filter(pk=company_id).update(closed_at=timezone.now() - dt.timedelta(days=89))
    assert cancellation.purge_closed_companies() == 0
    assert Company.objects.filter(pk=company_id).exists()

    Company.objects.filter(pk=company_id).update(closed_at=timezone.now() - dt.timedelta(days=90))
    assert cancellation.purge_closed_companies() == 1
    assert not Company.objects.filter(pk=company_id).exists()
    assert not User.objects.filter(company_id=company_id).exists()
    with company_context(company_id):
        assert not Customer.objects.filter(company_id=company_id).exists()
        assert not Payment.objects.filter(company_id=company_id).exists()
