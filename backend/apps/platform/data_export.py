"""D8 (decision 132): the company's whole information in one Excel file, a
sheet per table, so a financiera leaves (or just keeps a copy) with
everything it registered: clients, loans, installments, payments, cash,
collections, expenses, books, documents, messages and people. The terms
promise this copy during the 90 days after a cancellation.

Runs inside the company's tenant context. Every concrete field goes out as
text; a foreign key as the id of the related row (its own sheet has it).
"""

from __future__ import annotations

import datetime as dt
import io
from decimal import Decimal

from django.db import models
from django.utils import timezone

from apps.accounts.models import Invitation, Membership, Role, User
from apps.audit.models import DomainEvent
from apps.cash.models import CashDifference, CashSession, RouteSession
from apps.collections.models import (
    CollectionAction,
    CollectionTask,
    LegalCase,
    LegalCost,
    LegalProfessional,
    PaymentPromise,
)
from apps.core.models import Branch, Company, Holiday
from apps.customers.models import (
    Business,
    Collateral,
    Customer,
    Employment,
    Reference,
    Relative,
)
from apps.documents.models import GeneratedDocument
from apps.expenses.models import Expense, ExpenseCategory, Payable, PayablePayment
from apps.ledger.models import (
    AccountingPeriod,
    BankReconciliation,
    FinancialAccount,
    FiscalDocument,
    JournalEntry,
    JournalLine,
    LedgerAccount,
)
from apps.lending.models import (
    Application,
    Disbursement,
    Installment,
    LateFeePolicy,
    Loan,
    LoanCharge,
    LoanCollateral,
    LoanGuarantor,
    LoanProduct,
    LoanTerms,
    ProductCharge,
)
from apps.notifications.models import OutboundMessage
from apps.payments.models import LateFeeWaiver, Payment, PaymentAllocation, PaymentReversal
from apps.reports.models import CommissionLine, CommissionRule, CommissionRun, CustomerScore

# (sheet name ≤ 31 chars, model). The order reads like the business.
SHEETS: list[tuple[str, type[models.Model]]] = [
    ("Sucursales", Branch),
    ("Feriados", Holiday),
    ("Roles", Role),
    ("Membresías", Membership),
    ("Invitaciones", Invitation),
    ("Productos", LoanProduct),
    ("Mora de productos", LateFeePolicy),
    ("Cargos de productos", ProductCharge),
    ("Clientes", Customer),
    ("Empleos", Employment),
    ("Negocios", Business),
    ("Familiares", Relative),
    ("Referencias", Reference),
    ("Garantías", Collateral),
    ("Solicitudes", Application),
    ("Préstamos", Loan),
    ("Condiciones", LoanTerms),
    ("Cuotas", Installment),
    ("Cargos de préstamos", LoanCharge),
    ("Garantes", LoanGuarantor),
    ("Garantías de préstamos", LoanCollateral),
    ("Desembolsos", Disbursement),
    ("Pagos", Payment),
    ("Aplicación de pagos", PaymentAllocation),
    ("Reversos", PaymentReversal),
    ("Condonaciones", LateFeeWaiver),
    ("Cajas", CashSession),
    ("Rutas", RouteSession),
    ("Diferencias de caja", CashDifference),
    ("Gestiones", CollectionAction),
    ("Promesas", PaymentPromise),
    ("Seguimientos", CollectionTask),
    ("Profesionales legales", LegalProfessional),
    ("Casos legales", LegalCase),
    ("Costas", LegalCost),
    ("Categorías de gasto", ExpenseCategory),
    ("Gastos e ingresos", Expense),
    ("Cuentas por pagar", Payable),
    ("Pagos a proveedores", PayablePayment),
    ("Cuentas y bancos", FinancialAccount),
    ("Plan de cuentas", LedgerAccount),
    ("Asientos", JournalEntry),
    ("Líneas de asientos", JournalLine),
    ("Períodos", AccountingPeriod),
    ("Conciliaciones", BankReconciliation),
    ("Comprobantes fiscales", FiscalDocument),
    ("Documentos emitidos", GeneratedDocument),
    ("Mensajes enviados", OutboundMessage),
    ("Clasificación", CustomerScore),
    ("Reglas de comisión", CommissionRule),
    ("Corridas de comisión", CommissionRun),
    ("Comisiones", CommissionLine),
    ("Auditoría", DomainEvent),
]


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, dt.datetime):
        return timezone.localtime(value).strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, (dict, list)):
        return str(value)
    if isinstance(value, bool):
        return "sí" if value else "no"
    return str(value)


def _fields(model: type[models.Model]) -> list[models.Field]:
    return [
        f
        for f in model._meta.concrete_fields
        if f.name not in ("company",) and not f.name.startswith("_")
    ]


def _users_sheet(company_id) -> tuple[list[str], list[list[str]]]:
    header = ["id", "correo", "nombre", "apellido", "activo", "último acceso"]
    rows = [
        [str(u.id), u.email, u.first_name, u.last_name, _cell(u.is_active), _cell(u.last_login)]
        for u in User.objects.filter(company_id=company_id).order_by("email")
    ]
    return header, rows


def export_company(company: Company) -> bytes:
    """The workbook, in memory. Call inside ``company_context(company.id)``."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font

    book = Workbook(write_only=True)
    readme = book.create_sheet("Léeme")
    readme.append([f"Copia completa de {company.legal_name}"])
    readme.append([f"Generada el {timezone.localtime():%d/%m/%Y %H:%M} por Cobradía"])
    readme.append([])
    readme.append(
        [
            "Cada hoja es una tabla. Las columnas que terminan en «_id» apuntan a la fila "
            "con ese id en la hoja correspondiente. Los montos van en pesos con dos "
            "decimales; las fechas, como AAAA-MM-DD."
        ]
    )
    readme.append(
        [
            "Los archivos adjuntos (fotos de cédulas, documentos de expedientes, "
            "comprobantes) no van en este libro: pídelos a hola@cobradia.com."
        ]
    )
    bold = Font(bold=True)

    def heading(sheet, names: list[str]) -> None:
        cells = []
        for name in names:
            cell = WriteOnlyCell(sheet, value=name)
            cell.font = bold
            cells.append(cell)
        sheet.append(cells)

    header, rows = _users_sheet(company.id)
    sheet = book.create_sheet("Usuarios")
    heading(sheet, header)
    for row in rows:
        sheet.append(row)

    for name, model in SHEETS:
        fields = _fields(model)
        sheet = book.create_sheet(name[:31])
        heading(sheet, [f.column for f in fields])
        queryset = model.objects.filter(company_id=company.id).order_by("pk")
        for values in queryset.values_list(*[f.attname for f in fields]).iterator(chunk_size=2000):
            sheet.append([_cell(v) for v in values])

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
