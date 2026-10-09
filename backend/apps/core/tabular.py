"""Rows of a migration file (F16, decision 107): CSV or Excel, one header row.

Headers are matched without accents nor case («Teléfono» = «telefono»);
every value comes back as stripped text (Excel dates as YYYY-MM-DD,
whole numbers without «.0»), so each importer validates one format.
"""

import csv
import datetime as dt
import unicodedata
from pathlib import Path


def header(text: str) -> str:
    plain = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return plain.strip().lower().replace(" ", "_")


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_rows(path: str) -> list[dict[str, str]]:
    """The data rows as {header: text}, skipping fully empty ones."""
    if Path(path).suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return []
        keys = [header(cell) for cell in values[0]]
        raw = [dict(zip(keys, (_text(cell) for cell in row), strict=False)) for row in values[1:]]
    else:
        with open(path, encoding="utf-8-sig", newline="") as handle:
            raw = [
                {header(key): _text(value) for key, value in row.items() if key}
                for row in csv.DictReader(handle)
            ]
    return [row for row in raw if any(row.values())]
