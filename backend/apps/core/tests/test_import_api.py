"""D1 (decision 131): the portfolio import from the app."""

import datetime as dt
import io

from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from openpyxl import load_workbook

from apps.core.tenancy import company_context
from apps.customers.models import Customer
from apps.lending.models import Loan

SAN = "Personal San (semanal)"


def _upload(client, kind: str, name: str, content: bytes, apply: bool = False):
    file = SimpleUploadedFile(name, content)
    data = {"archivo": file, "aplicar": "true" if apply else "false"}
    return client.post(f"/api/v1/importacion/{kind}/", data, format="multipart")


def test_templates_download_for_the_manager_only(tenant) -> None:
    manager = tenant.client("gerente")
    response = manager.get("/api/v1/importacion/plantillas/clientes/")
    assert response.status_code == 200
    assert response["Content-Disposition"] == 'attachment; filename="clientes.xlsx"'
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Datos", "Instrucciones"]
    assert [c.value for c in book["Datos"][1]][:3] == ["tipo", "tipo_documento", "documento"]
    assert manager.get("/api/v1/importacion/plantillas/prestamos/").status_code == 200
    assert manager.get("/api/v1/importacion/plantillas/otra/").status_code == 404
    assert (
        tenant.client("cajero").get("/api/v1/importacion/plantillas/clientes/").status_code == 403
    )


def test_customers_rehearse_then_apply_from_the_app(tenant) -> None:
    manager = tenant.client("gerente")
    header = "tipo,tipo_documento,documento,nombres,apellidos,telefono,sucursal\n"
    good = "persona,cedula,002-1234567-2,Pedro,Santana,8095550102,01\n"
    bad = "persona,cedula,999-9999999-9,Rosa,Mejía,,01\n"
    with company_context(tenant.company.id):
        before = Customer.objects.count()

    report = _upload(manager, "clientes", "clientes.csv", (header + good + bad).encode()).json()
    assert (report["created"], report["errors"], report["applied"]) == (1, 1, False)
    assert [row["status"] for row in report["rows"]] == ["created", "error"]
    assert "documento" in report["rows"][1]["detail"]
    assert report["summary"].startswith("Nada se guardó")

    # An apply with an error saves nothing either.
    report = _upload(manager, "clientes", "clientes.csv", (header + good + bad).encode(), True)
    assert report.json()["applied"] is False
    with company_context(tenant.company.id):
        assert Customer.objects.count() == before

    report = _upload(manager, "clientes", "clientes.csv", (header + good).encode(), True).json()
    assert report["applied"] is True and report["summary"].startswith("Importación aplicada")
    with company_context(tenant.company.id):
        assert Customer.objects.count() == before + 1
    # Loading the corrected file again skips what already exists.
    again = _upload(manager, "clientes", "clientes.csv", (header + good).encode(), True).json()
    assert (again["created"], again["existing"]) == (0, 1)


def test_loans_import_points_at_the_missing_customer(tenant) -> None:
    manager = tenant.client("gerente")
    start = timezone.localdate() - dt.timedelta(days=14)
    header = (
        "documento,producto,monto,tasa,frecuencia,cuotas,fecha_desembolso,primera_cuota,"
        "numero_anterior,cobrador,saldo_capital,pagos\n"
    )
    good = (
        f"001-1234567-4,{SAN},1000,0.10,semanal,4,{start},,A-1,cobrador@uno.do,750.00,"
        f"{start + dt.timedelta(days=7)}:350\n"
    )
    missing = f"999-9999999-9,{SAN},1000,0.10,semanal,4,{start},,A-2,,1000.00,\n"
    report = _upload(manager, "prestamos", "prestamos.csv", (header + good + missing).encode())
    rows = report.json()["rows"]
    assert rows[0]["status"] == "created", rows[0] and "A-1 →" in rows[0]["detail"]
    assert rows[1]["status"] == "error" and "carga primero los clientes" in rows[1]["detail"]
    with company_context(tenant.company.id):
        assert not Loan.objects.exists()

    done = _upload(manager, "prestamos", "prestamos.csv", (header + good).encode(), True).json()
    assert done["applied"] is True
    with company_context(tenant.company.id):
        loan = Loan.objects.get(legacy_number="A-1")
        assert loan.collector == tenant.users["cobrador"]


def test_whole_file_problems_are_one_message(tenant) -> None:
    manager = tenant.client("gerente")
    assert manager.post("/api/v1/importacion/clientes/", {}, format="multipart").status_code == 400
    wrong = _upload(manager, "clientes", "clientes.pdf", b"%PDF")
    assert wrong.status_code == 400 and "Excel" in wrong.json()["archivo"]
    empty = _upload(manager, "clientes", "clientes.csv", b"tipo,documento\n")
    assert empty.status_code == 400 and "filas" in empty.json()["archivo"]
    assert _upload(tenant.client("cajero"), "clientes", "c.csv", b"a,b\n1,2").status_code == 403
