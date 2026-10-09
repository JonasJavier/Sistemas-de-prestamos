"""One ``Table``, three faces (decision 087): JSON for the screen, a PDF
on letter paper and an Excel sheet with real numbers, all with the same
summary, filters line, totals and note (doc 09 §1)."""

import datetime as dt
import io
from decimal import Decimal

from django.utils import timezone

from apps.reports.catalog import Table

KINDS_RIGHT = {"money", "int", "pct"}


def cell(value, kind: str) -> str:
    """Display text, as the screen shows it."""
    if value is None or value == "":
        return "" if kind == "text" else "—"
    if kind == "money":
        return f"{Decimal(value):,.2f}"
    if kind == "int":
        return f"{int(value):,}"
    if kind == "pct":
        text = f"{Decimal(value) * 100:,.1f}".rstrip("0").rstrip(".")
        return f"{text} %"
    if kind == "date":
        return value.strftime("%d/%m/%Y") if isinstance(value, dt.date) else str(value)
    return str(value)


def money_or_cell(value, kind: str) -> str:
    text = cell(value, kind)
    return f"RD$ {text}" if kind == "money" and text != "—" else text


def as_json(table: Table, *, limit: int | None = None) -> dict:
    rows = table.rows if limit is None else table.rows[:limit]
    return {
        "title": table.title,
        "filters_line": table.filters_line,
        "columns": [
            {"key": c.key, "label": c.label, "kind": c.kind, "personal": c.personal}
            for c in table.columns
        ],
        "rows": rows,
        "count": len(table.rows),
        "truncated": len(rows) < len(table.rows),
        "totals": table.totals,
        "summary": table.summary,
        "note": table.note,
    }


def to_pdf(table: Table, *, company, user) -> bytes:
    from apps.documents.services import _monogram, render_pdf

    now = timezone.localtime()
    issuer = (user.get_full_name() or user.email).replace('"', "'").replace("\\", "/")
    totals = table.totals
    return render_pdf(
        "reports/report.html",
        {
            "company": company,
            "monogram": _monogram(company),
            "branch": {"name": ""},
            "folio": "",
            "issuer_label": issuer,
            "issued_at": now,
            "legal_pending": False,
            "table": table,
            "summary": [
                {**item, "text": money_or_cell(item.get("value"), item.get("kind", "text"))}
                for item in table.summary
            ],
            "head": [(c.label, c.kind in KINDS_RIGHT) for c in table.columns],
            "body": [
                [(cell(row.get(c.key), c.kind), c.kind in KINDS_RIGHT) for c in table.columns]
                for row in table.rows
            ],
            "foot": [
                (cell(totals[c.key], c.kind) if c.key in totals else "", c.kind in KINDS_RIGHT)
                for c in table.columns
            ]
            if totals
            else None,
        },
    )


def to_xlsx(table: Table, *, company, user) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = table.title[:31]
    bold = Font(bold=True)
    sheet.append([company.legal_name])
    sheet["A1"].font = Font(bold=True, size=13, color="15402F")
    sheet.append([table.title])
    sheet["A2"].font = Font(bold=True, size=12)
    sheet.append([table.filters_line])
    who = user.get_full_name() or user.email
    sheet.append([f"Generado por {who} el {timezone.localtime():%d/%m/%Y %H:%M}"])
    sheet.append([])
    formats = {"money": "#,##0.00", "int": "#,##0", "pct": "0.0%", "date": "dd/mm/yyyy"}

    def put(row_cells: list, kinds: list[str]) -> None:
        sheet.append(row_cells)
        for index, kind in enumerate(kinds, start=1):
            if kind in formats:
                sheet.cell(row=sheet.max_row, column=index).number_format = formats[kind]

    for item in table.summary:
        kind = item.get("kind", "text")
        value = item.get("value")
        put([item["label"], _number(value, kind), item.get("hint", "")], ["text", kind, "text"])
    if table.summary:
        sheet.append([])
    sheet.append([c.label for c in table.columns])
    fill = PatternFill("solid", fgColor="EEF2EF")
    for index in range(1, len(table.columns) + 1):
        header = sheet.cell(row=sheet.max_row, column=index)
        header.font = bold
        header.fill = fill
        header.alignment = Alignment(wrap_text=True)
    header_row = sheet.max_row
    kinds = [c.kind for c in table.columns]
    for row in table.rows:
        put([_number(row.get(c.key), c.kind) for c in table.columns], kinds)
    totals = table.totals
    if totals:
        put(
            ["Total"] + [_number(totals.get(c.key), c.kind) for c in table.columns[1:]],
            ["text"] + kinds[1:],
        )
        for index in range(1, len(table.columns) + 1):
            sheet.cell(row=sheet.max_row, column=index).font = bold
    if table.note:
        sheet.append([])
        sheet.append([table.note])
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    for index, column in enumerate(table.columns, start=1):
        width = 14 if column.kind in KINDS_RIGHT | {"date"} else 24
        sheet.column_dimensions[get_column_letter(index)].width = width
    # Text that starts with «=» (a company's or a customer's name, a note)
    # stays text: this sheet never carries formulas Excel would run (102).
    for line in sheet.iter_rows():
        for cell in line:
            if cell.data_type == "f":
                cell.data_type = "s"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _number(value, kind: str):
    """Excel gets numbers and dates, not text, so it can add them up."""
    if value is None or value == "":
        return None
    if kind in {"money", "pct"}:
        return float(Decimal(value))  # display only; the exact figure is the PDF and screen
    if kind == "int":
        return int(value)
    return value
