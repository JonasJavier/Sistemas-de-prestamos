"""Phase 7 · P64: the report catalog by role, the coherence rule (screen,
PDF and Excel from the same Table) and the background exports (087)."""

import datetime as dt
import io
from decimal import Decimal

from django.utils import timezone
from openpyxl import load_workbook

from apps.audit.models import SensitiveAccessLog
from apps.core.tenancy import company_context
from apps.payments.tasks import run_nightly
from apps.payments.tests.test_payments import make_loan, pay
from apps.reports import exports, indicators
from apps.reports.models import ExportJob

D = Decimal


def today() -> dt.date:
    return timezone.localdate()


def keys(tenant, role):
    return [r["key"] for r in tenant.client(role).get("/api/v1/reportes/").json()["reports"]]


def export(tenant, role, key, fmt="xlsx", capture=None, **filters):
    with capture(execute=True):
        response = tenant.client(role).post(
            f"/api/v1/reportes/{key}/exportar/",
            {"format": fmt, "filters": filters},
            format="json",
        )
    return response


def test_each_role_sees_its_reports(tenant) -> None:
    manager = keys(tenant, "gerente")
    assert manager == [
        "cartera",
        "atrasos",
        "cobros",
        "desembolsos",
        "caja",
        "gestiones",
        "resultados",
        "comisiones",
        "clasificacion",
    ]
    assert "resultados" not in keys(tenant, "supervisor")
    assert "comisiones" not in keys(tenant, "supervisor")
    assert keys(tenant, "contabilidad") == manager[:-1]  # doc 07: P62 is G, Su, O and Au
    assert keys(tenant, "oficial") == ["atrasos", "clasificacion"]
    assert keys(tenant, "cobrador") == []
    assert keys(tenant, "cajero") == []
    assert tenant.client("oficial").get("/api/v1/reportes/cobros/datos/").status_code == 404
    assert tenant.client("supervisor").get("/api/v1/reportes/resultados/datos/").status_code == 404


def test_the_reports_match_the_indicators(tenant) -> None:
    loan = make_loan(tenant, first_due_in=-10)
    make_loan(tenant, first_due_in=3)
    pay(tenant, loan, "500")
    run_nightly_today(tenant)
    client = tenant.client("gerente")
    with company_context(tenant.company.id):
        collected = indicators.collected(today())
        portfolio = indicators.portfolio(today())
        month = indicators.disbursed(today().year, today().month)

    payments = client.get("/api/v1/reportes/cobros/datos/").json()
    assert D(payments["totals"]["amount"]) == collected == D("500.00")
    parts = sum(D(payments["totals"][k]) for k in ("principal", "interest", "late_fee", "charges"))
    assert parts == D("500.00")  # the components add up to what was applied

    cartera = client.get("/api/v1/reportes/cartera/datos/").json()
    assert D(cartera["totals"]["principal"]) == portfolio.principal
    assert cartera["summary"][0] == {
        "label": "Cartera activa",
        "value": str(portfolio.principal),
        "kind": "money",
    }

    disbursed = client.get("/api/v1/reportes/desembolsos/datos/").json()
    assert D(disbursed["totals"]["gross"]) == month["total"]

    arrears = client.get("/api/v1/reportes/atrasos/datos/").json()
    assert arrears["count"] == 1 and arrears["rows"][0]["loan"] == loan.number
    assert "Corte al" in arrears["filters_line"]

    assert client.get("/api/v1/reportes/cobros/datos/?desde=2026-13-01").status_code == 400
    assert client.get("/api/v1/reportes/cartera/datos/?corte=color").status_code == 400


def run_nightly_today(tenant):
    with company_context(tenant.company.id):
        run_nightly(tenant.company.id, today())


def test_a_manager_exports_personal_data_and_it_is_logged(
    tenant, django_capture_on_commit_callbacks
) -> None:
    make_loan(tenant, first_due_in=-10)
    run_nightly_today(tenant)
    response = export(tenant, "gerente", "atrasos", capture=django_capture_on_commit_callbacks)
    assert response.status_code == 202, response.content
    job_id = response.json()["id"]
    mine = tenant.client("gerente").get("/api/v1/exportaciones/").json()
    assert mine[0]["status"] == "done" and mine[0]["personal"] is True and mine[0]["rows"] == 1
    file = tenant.client("gerente").get(f"/api/v1/exportaciones/{job_id}/archivo/")
    assert file.status_code == 200
    sheet = load_workbook(io.BytesIO(file.content)).active
    values = [cell for row in sheet.iter_rows(values_only=True) for cell in row if cell]
    assert "Cliente" in values and "María Gómez" in values
    with company_context(tenant.company.id):
        assert SensitiveAccessLog.objects.filter(
            action=SensitiveAccessLog.Action.EXPORT, target_id=job_id
        ).exists()
    # Only its owner downloads it.
    assert (
        tenant.client("gerente2").get(f"/api/v1/exportaciones/{job_id}/archivo/").status_code == 404
    )


def test_a_supervisor_exports_without_personal_data(
    tenant, django_capture_on_commit_callbacks
) -> None:
    make_loan(tenant, first_due_in=-10)
    run_nightly_today(tenant)
    response = export(tenant, "supervisor", "atrasos", capture=django_capture_on_commit_callbacks)
    assert response.status_code == 202
    assert response.json()["personal"] is False
    job_id = response.json()["id"]
    file = tenant.client("supervisor").get(f"/api/v1/exportaciones/{job_id}/archivo/")
    sheet = load_workbook(io.BytesIO(file.content)).active
    values = [cell for row in sheet.iter_rows(values_only=True) for cell in row if cell]
    assert "Cliente" not in values and "María Gómez" not in values
    assert any("Sin datos personales" in str(v) for v in values)


def test_pdf_export_and_who_may_export(tenant, django_capture_on_commit_callbacks) -> None:
    make_loan(tenant, first_due_in=3)
    response = export(
        tenant, "contabilidad", "cartera", fmt="pdf", capture=django_capture_on_commit_callbacks
    )
    assert response.status_code == 202
    file = tenant.client("contabilidad").get(
        f"/api/v1/exportaciones/{response.json()['id']}/archivo/"
    )
    assert file["Content-Type"] == "application/pdf" and file.content.startswith(b"%PDF")

    officer = export(tenant, "oficial", "atrasos", capture=django_capture_on_commit_callbacks)
    assert officer.status_code == 403  # reads the report, does not export it
    collector = export(tenant, "cobrador", "cartera", capture=django_capture_on_commit_callbacks)
    assert collector.status_code == 403
    bad = export(
        tenant,
        "gerente",
        "cobros",
        capture=django_capture_on_commit_callbacks,
        desde="2026-10-09",
        hasta="2026-10-01",
    )
    assert bad.status_code == 400


def test_old_export_files_expire(tenant, django_capture_on_commit_callbacks) -> None:
    response = export(tenant, "gerente", "caja", capture=django_capture_on_commit_callbacks)
    with company_context(tenant.company.id):
        job = ExportJob.objects.get(pk=response.json()["id"])
        ExportJob.objects.filter(pk=job.pk).update(
            finished_at=timezone.now() - dt.timedelta(days=8)
        )
        assert exports.purge(today()) == 1
        job.refresh_from_db()
        assert job.status == ExportJob.Status.EXPIRED and not job.file
    gone = tenant.client("gerente").get(f"/api/v1/exportaciones/{job.id}/archivo/")
    assert gone.status_code == 404


def test_reports_and_exports_never_cross_companies(
    tenant, other_tenant, django_capture_on_commit_callbacks
) -> None:
    make_loan(tenant, first_due_in=-10)
    run_nightly_today(tenant)
    response = export(tenant, "gerente", "atrasos", capture=django_capture_on_commit_callbacks)
    job_id = response.json()["id"]
    theirs = other_tenant.client("gerente")
    assert theirs.get(f"/api/v1/exportaciones/{job_id}/archivo/").status_code == 404
    assert theirs.get("/api/v1/exportaciones/").json() == []
    assert theirs.get("/api/v1/reportes/atrasos/datos/").json()["count"] == 0
    picked = theirs.get(f"/api/v1/reportes/cartera/datos/?sucursal={tenant.branch.id}")
    assert picked.status_code == 404
